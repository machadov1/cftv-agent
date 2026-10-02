"""Planilha de controle das RITMs (skill operador-cftv, seção 14).

Regras: nunca recriar o arquivo; primeira linha com A vazia (a partir da 2, sem usar max_row);
copiar o estilo da linha de cima; manter a fórmula da coluna F. Faz backup antes de salvar.
Em dry-run/mock devolve só a linha que seria gravada.
"""
import re
import shutil
from datetime import datetime

import openpyxl

from backend.config import config, ROOT
from backend.ritm import _ANALISE, PENDENCIAS

SHEET = "Controle"
BACKUP_DIR = ROOT / "data" / "backups"
COLUNAS = ["INC", "REQ", "RITM", "ABERTURA", "DIAS", "FECHAMENTO", "SITUACAO", "UNIDADE",
           "METALICOS", "DESCRICAO", "APROVADA", "PENDENCIA"]


def is_metalicos(text: str) -> bool:
    t = (text or "").upper()
    return bool(re.search(r"SUC(?![A-Z])|SUCATA|MET[ÁA]LICOS", t))


def default_resumo(causa: str, pendencia: str) -> str:
    dep = _ANALISE.get(pendencia, pendencia)
    return f"{causa} - depende de {dep}" if pendencia else causa


def build_row(inc_number: str, req: str, ritm: str, unidade: str, resumo: str, pendencia: str,
              metalicos: bool, data: datetime | None = None) -> dict:
    if pendencia not in PENDENCIAS:
        raise ValueError(f"Pendência inválida ({', '.join(PENDENCIAS)})")
    d = (data or datetime.now()).replace(hour=0, minute=0, second=0, microsecond=0)
    return {
        "A": inc_number, "B": req, "C": ritm, "D": d, "E": 20,
        # F: fórmula (preenchida abaixo com o número da linha, se a linha ainda não a tiver)
        "G": "Aberta", "H": re.sub(r"\s*\(LORA\)", "", unidade or "").strip(), "I": "Sim" if metalicos else "Não",
        "J": resumo, "K": None, "L": f"Pendência de {pendencia}",
    }


def _first_empty_row(ws) -> int:
    r = 2
    while ws.cell(r, 1).value not in (None, ""):
        r += 1
    return r


def append_row(row: dict, path: str | None = None, *, apply: bool) -> dict:
    """Localiza a linha e (se apply) grava. Retorna o que foi/seria gravado."""
    path = path or config.CONTROLE_XLSX
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]

    # Já lançada? (coluna C = RITM)
    for r in range(2, _first_empty_row(ws)):
        if ws.cell(r, 3).value == row["C"]:
            return {"ok": False, "duplicada": True, "linha": r, "erro": f"{row['C']} já está na planilha (linha {r})."}

    n = _first_empty_row(ws)
    out = {"ok": True, "linha": n, "arquivo": path, "valores": {**row, "F": ws.cell(n, 6).value or f'=IF(D{n}="","",D{n}+E{n})'},
           "gravado": False}
    out["valores"]["D"] = row["D"].strftime("%d/%m/%Y")
    if not apply:
        return out

    for col_idx, letter in enumerate("ABCDEFGHIJKL", start=1):
        c = ws.cell(n, col_idx)
        prev = ws.cell(n - 1, col_idx)
        if prev.has_style:
            c._style = prev._style
        if letter == "F":
            if c.value in (None, ""):
                c.value = f'=IF(D{n}="","",D{n}+E{n})'
            continue
        c.value = row[letter]
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, BACKUP_DIR / f"controle_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    try:
        wb.save(path)
    except PermissionError:
        out.update(ok=False, erro="Arquivo aberto no Excel: feche a planilha e tente de novo (nada foi gravado).")
        return out
    out["gravado"] = True
    out["aviso"] = "Fórmulas recalculam quando o Excel abrir o arquivo."
    return out


_leitura: dict = {"mtime": None, "rows": []}


def read_rows(path: str | None = None) -> tuple[list[dict], str | None]:
    """RITMs lançadas (só leitura, para as métricas). (linhas, aviso). Fechamento = coluna F ou D + E dias."""
    from datetime import timedelta
    from pathlib import Path
    p = Path(path or config.CONTROLE_XLSX)
    try:
        mtime = p.stat().st_mtime
    except OSError:
        return [], "Planilha de controle não encontrada."
    if _leitura["mtime"] == mtime:
        return _leitura["rows"], None
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[SHEET]
        rows = []
        for v in ws.iter_rows(min_row=2, max_col=12, values_only=True):
            if not v or not v[0]:
                continue
            abertura = v[3] if isinstance(v[3], datetime) else None
            dias = v[4] if isinstance(v[4], (int, float)) else 20
            fecha = v[5] if isinstance(v[5], datetime) else (abertura + timedelta(days=dias) if abertura else None)
            rows.append({"inc": str(v[0]).strip(), "req": v[1], "ritm": v[2], "abertura": abertura, "fechamento": fecha,
                         "situacao": (v[6] or "").strip(), "unidade": (v[7] or "").strip(),
                         "descricao": (v[9] or "").strip(),
                         "pendencia": re.sub(r"^Pend[êe]ncia de\s*", "", (v[11] or "").strip(), flags=re.I)})
        wb.close()
    except Exception as e:  # noqa: BLE001 (arquivo travado/corrompido: as métricas seguem sem as RITMs)
        return _leitura["rows"], f"Não consegui ler a planilha: {e}"
    _leitura.update(mtime=mtime, rows=rows)
    return rows, None
