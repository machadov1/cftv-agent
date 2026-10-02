import datetime as dt

import openpyxl
import pytest

from backend import teams, controle


def test_saudacao_e_nome():
    assert teams.saudacao(dt.datetime(2026, 9, 29, 9)) == "Bom dia"
    assert teams.saudacao(dt.datetime(2026, 9, 29, 14)) == "Boa tarde"
    assert teams.saudacao(dt.datetime(2026, 9, 29, 20)) == "Boa noite"
    assert teams.primeiro_nome("Costa, Ericon Sampaio") == "Ericon"


def test_mensagem_e_url_nao_enviam():
    m = teams.build_message("INC1", "Ericon", "O reparo depende da montagem de andaime.", "RITM1",
                            "Câmera PIR400 sem conexão", dt.datetime(2026, 9, 29, 9))
    assert m == ("Bom dia, Ericon! Tudo bem?\n\n"
                 "Passando pra te atualizar sobre o incidente INC1, câmera PIR400 sem conexão.\n\n"
                 "O reparo depende da montagem de andaime.\n\n"
                 "Para acompanhamento das ações, foi aberta a requisição RITM1.\n\n"
                 "Qualquer dúvida, fico à disposição!")
    u = teams.build_url("a@b.com", m)
    assert u.startswith("https://teams.cloud.microsoft/l/chat/0/0?users=a%40b.com&message=")
    assert "\n" not in u


def _planilha(tmp_path):
    p = tmp_path / "c.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controle"
    ws.append(["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"])
    ws.append(["INC1", "REQ1", "RITM0", dt.datetime(2026, 9, 1), 20, "=D2+E2", "Aberta", "Resende", "Não", "x", None,
               "Pendência de PEMT"])
    for r in range(3, 8):  # linhas com fórmula em F mas A vazia (max_row engana)
        ws.cell(r, 6).value = f'=IF(D{r}="","",D{r}+E{r})'
    wb.save(p)
    return str(p)


def test_planilha_linha_certa_e_duplicidade(tmp_path, monkeypatch):
    monkeypatch.setattr(controle, "BACKUP_DIR", tmp_path / "bk")
    p = _planilha(tmp_path)
    row = controle.build_row("INC9", "REQ9", "RITM9", "Resende", "Câmera RES024 sem conexão - depende de PEMT", "PEMT", False)
    prev = controle.append_row(row, p, apply=False)
    assert prev["linha"] == 3 and not prev["gravado"]
    assert openpyxl.load_workbook(p)["Controle"].cell(3, 1).value is None  # preview não grava

    res = controle.append_row(row, p, apply=True)
    assert res["gravado"]
    ws = openpyxl.load_workbook(p)["Controle"]
    assert [ws.cell(3, c).value for c in (1, 2, 3, 5, 7, 8, 9, 12)] == [
        "INC9", "REQ9", "RITM9", 20, "Aberta", "Resende", "Não", "Pendência de PEMT"]
    assert ws.cell(3, 6).value.startswith("=IF(D3")
    assert list((tmp_path / "bk").glob("*.xlsx"))  # backup feito

    assert controle.append_row(row, p, apply=True)["duplicada"]


def test_metalicos_e_pendencia_invalida():
    assert controle.is_metalicos("Câmeras 03SUC/10SUC") and not controle.is_metalicos("Câmera RES024")
    with pytest.raises(ValueError):
        controle.build_row("I", "R", "T", "U", "j", "Outra", False)
